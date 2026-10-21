"""Tests for the right-to-left repair pass applied to a generated DOCX."""

import os
import tempfile

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from pdf_editor.docx_rtl import apply_rtl


def make_docx(paragraphs: list[str], path: str) -> None:
    """Writes a DOCX containing one run per supplied paragraph."""
    document = Document()
    for text in paragraphs:
        document.add_paragraph(text)
    document.save(path)


def docx_xml(path: str) -> str:
    """Returns the document body XML as a string."""
    document = Document(path)
    return document.element.body.xml


def temp_docx() -> str:
    with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as f:
        return f.name


def _pPr(paragraph):
    """Returns the paragraph's pPr, creating it if the paragraph has none."""
    existing = paragraph._element.find(qn("w:pPr"))
    if existing is None:
        existing = OxmlElement("w:pPr")
        paragraph._element.insert(0, existing)
    return existing


def set_geometry(path: str, jc: str | None = None, **indents: str) -> None:
    """
    Writes w:jc and w:ind onto the first paragraph, the way pdf2docx does.

    Args:
        path (str): DOCX to modify in place.
        jc (str): Value for w:jc, e.g. "left".
        **indents: w:ind attributes without their prefix, e.g. left="4762".
    """
    document = Document(path)
    pPr = _pPr(document.paragraphs[0])

    if jc is not None:
        alignment = OxmlElement("w:jc")
        alignment.set(qn("w:val"), jc)
        pPr.append(alignment)

    if indents:
        indent = OxmlElement("w:ind")
        for name, value in indents.items():
            indent.set(qn(f"w:{name}"), value)
        pPr.append(indent)

    document.save(path)


def alignment_of(path: str) -> str | None:
    """Returns the first paragraph's w:jc value."""
    pPr = Document(path).paragraphs[0]._element.find(qn("w:pPr"))
    alignment = pPr.find(qn("w:jc"))
    return alignment.get(qn("w:val")) if alignment is not None else None


def indents_of(path: str) -> dict[str, str]:
    """Returns the first paragraph's w:ind attributes, keyed without prefixes."""
    pPr = Document(path).paragraphs[0]._element.find(qn("w:pPr"))
    indent = pPr.find(qn("w:ind"))
    if indent is None:
        return {}
    return {
        name.split("}")[-1]: value
        for name, value in indent.attrib.items()
    }


class TestApplyRtl:
    def test_marks_rtl_paragraphs_and_runs(self):
        path = temp_docx()
        try:
            make_docx(["שלום עולם"], path)
            stats = apply_rtl(path)

            assert stats["paragraphs"] == 1
            assert stats["runs"] == 1

            xml = docx_xml(path)
            assert "<w:bidi" in xml
            assert "<w:rtl" in xml
        finally:
            os.unlink(path)

    def test_leaves_ltr_documents_alone(self):
        path = temp_docx()
        try:
            make_docx(["Hello world", "Nothing to see here"], path)
            before = docx_xml(path)
            stats = apply_rtl(path)

            assert stats == {
                "paragraphs": 0,
                "runs": 0,
                "mirrored": 0,
                "text_fixed": 0,
                "text_unmatched": 0,
            }
            assert docx_xml(path) == before
        finally:
            os.unlink(path)

    def test_substitutes_text_from_the_correction_table(self):
        path = temp_docx()
        try:
            make_docx([":נושא המחקר"], path)
            stats = apply_rtl(path, {":נושא המחקר": "נושא המחקר:"})

            assert stats["text_fixed"] == 1
            assert Document(path).paragraphs[0].text == "נושא המחקר:"
        finally:
            os.unlink(path)

    def test_matches_on_collapsed_whitespace_and_keeps_own_padding(self):
        path = temp_docx()
        try:
            make_docx(["  :שם החוקר "], path)
            # The table entry carries different padding than the run does.
            stats = apply_rtl(path, {":שם החוקר": "שם החוקר:      "})

            assert stats["text_fixed"] == 1
            assert Document(path).paragraphs[0].text == "  שם החוקר: "
        finally:
            os.unlink(path)

    def test_unmatched_text_is_left_untouched(self):
        path = temp_docx()
        try:
            make_docx(["טקסט שאין לו התאמה"], path)
            stats = apply_rtl(path, {"משהו אחר": "אחר משהו"})

            assert stats["text_fixed"] == 0
            assert stats["text_unmatched"] == 1
            assert Document(path).paragraphs[0].text == "טקסט שאין לו התאמה"
        finally:
            os.unlink(path)

    def test_reaches_paragraphs_inside_tables(self):
        path = temp_docx()
        try:
            document = Document()
            table = document.add_table(rows=1, cols=2)
            table.cell(0, 0).text = ":מחלקה"
            table.cell(0, 1).text = "Department"
            document.save(path)

            stats = apply_rtl(path, {":מחלקה": "מחלקה:"})

            assert stats["text_fixed"] == 1
            assert Document(path).tables[0].cell(0, 0).text == "מחלקה:"
            # The Latin cell must be left as it was.
            assert Document(path).tables[0].cell(0, 1).text == "Department"
        finally:
            os.unlink(path)

    def test_result_is_still_a_readable_docx(self):
        path = temp_docx()
        try:
            make_docx(["שלום", "Hello", "עולם 2026"], path)
            apply_rtl(path)

            document = Document(path)
            assert [p.text for p in document.paragraphs] == ["שלום", "Hello", "עולם 2026"]
        finally:
            os.unlink(path)

    def test_run_properties_mirror_to_complex_script(self):
        path = temp_docx()
        try:
            document = Document()
            run = document.add_paragraph().add_run("שלום")
            run.bold = True
            run.font.size = 152400  # 12pt in EMU
            document.save(path)

            apply_rtl(path)

            rPr = Document(path).paragraphs[0].runs[0]._element.find(qn("w:rPr"))
            assert rPr.find(qn("w:rtl")) is not None
            assert rPr.find(qn("w:bCs")) is not None
            assert rPr.find(qn("w:szCs")) is not None
        finally:
            os.unlink(path)

    def test_is_idempotent_including_geometry(self):
        """Mirroring twice would undo itself, so it must happen only once."""
        path = temp_docx()
        try:
            make_docx(['שלום עולם'], path)
            set_geometry(path, jc='right', left='0', right='104')

            first = apply_rtl(path)
            after_first = docx_xml(path)
            second = apply_rtl(path)

            assert first['mirrored'] == 1
            assert second['mirrored'] == 0
            assert docx_xml(path) == after_first
        finally:
            os.unlink(path)


class TestMirrorGeometry:
    """
    w:jc left/right and w:ind left/right are the Transitional spellings of
    start/end, so they flip meaning once w:bidi is set. pdf2docx computes them
    for left-to-right layout; they have to be mirrored to survive the change.
    """

    def test_right_alignment_becomes_left(self):
        path = temp_docx()
        try:
            make_docx(['נושא המחקר:'], path)
            set_geometry(path, jc='right')

            apply_rtl(path)

            assert alignment_of(path) == 'left'
        finally:
            os.unlink(path)

    def test_left_alignment_becomes_right(self):
        path = temp_docx()
        try:
            make_docx(['מרכז רפואי רבין'], path)
            set_geometry(path, jc='left')

            apply_rtl(path)

            assert alignment_of(path) == 'right'
        finally:
            os.unlink(path)

    def test_centre_alignment_is_unchanged(self):
        path = temp_docx()
        try:
            make_docx(['טיפול בבקשה חדשה'], path)
            set_geometry(path, jc='center')

            apply_rtl(path)

            assert alignment_of(path) == 'center'
        finally:
            os.unlink(path)

    def test_indents_swap_sides(self):
        path = temp_docx()
        try:
            make_docx(['מרכז רפואי רבין'], path)
            set_geometry(path, jc='left', left='4762', right='4032')

            apply_rtl(path)

            assert indents_of(path) == {'left': '4032', 'right': '4762'}
        finally:
            os.unlink(path)

    def test_a_one_sided_indent_moves_to_the_other_side(self):
        """The absent side must end up absent, not carried over as a stale value."""
        path = temp_docx()
        try:
            make_docx(['נושא המחקר:'], path)
            set_geometry(path, jc='right', right='104')

            apply_rtl(path)

            assert indents_of(path) == {'left': '104'}
        finally:
            os.unlink(path)

    def test_first_line_indent_is_left_alone(self):
        """firstLine is measured from the start edge, so it needs no mirroring."""
        path = temp_docx()
        try:
            make_docx(['נושא המחקר:'], path)
            set_geometry(path, jc='right', left='200', firstLine='360')

            apply_rtl(path)

            assert indents_of(path)['firstLine'] == '360'
        finally:
            os.unlink(path)

    def test_section_direction_is_left_untouched(self):
        """
        Section-level bidi would make RTL the default for paragraphs this pass
        never mirrors, moving Latin text and numbers to the wrong side.
        """
        path = temp_docx()
        try:
            make_docx(['שלום עולם'], path)
            apply_rtl(path)

            sectPr = Document(path).element.body.find(qn('w:sectPr'))
            assert sectPr is not None
            assert sectPr.find(qn('w:bidi')) is None
        finally:
            os.unlink(path)

    def test_ltr_paragraphs_keep_their_geometry(self):
        path = temp_docx()
        try:
            make_docx(['Latin only'], path)
            set_geometry(path, jc='right', left='0', right='104')

            apply_rtl(path)

            assert alignment_of(path) == 'right'
            assert indents_of(path) == {'left': '0', 'right': '104'}
        finally:
            os.unlink(path)

    def test_is_idempotent(self):
        path = temp_docx()
        try:
            make_docx(["שלום עולם"], path)
            apply_rtl(path)
            first = docx_xml(path)
            apply_rtl(path)

            assert docx_xml(path) == first
        finally:
            os.unlink(path)
