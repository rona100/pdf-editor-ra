import io
import pytest
import shutil
import tempfile
import zipfile
import os
from pathlib import Path

import fitz
from docx import Document
from pypdf import PdfReader, PdfWriter

# Import the functions to test
from pdf_editor import pdf_operations
from pdf_editor.pdf_operations import (
    convert_docx_to_pdf,
    convert_pdf_to_docx,
    merge_pdfs,
    order_pdf_pages,
    rotate_pdf_pages,
    ats_pdf,
    sanitize_pdf,
    scrub_pdf,
)
from test_rtl import HEBREW_FONT, needs_hebrew_font

FIXTURES = Path(__file__).parent / 'fixtures'

needs_soffice = pytest.mark.skipif(
    shutil.which('soffice') is None and shutil.which('libreoffice') is None,
    reason='LibreOffice is not installed',
)

needs_tesseract = pytest.mark.skipif(
    shutil.which('tesseract') is None,
    reason='Tesseract is not installed',
)


class TestPDFOperations:
    """Test class for PDF rotation and merging operations."""

    def create_test_pdf(self, pages=3, content="Test Page"):
        """Create a simple test PDF with specified number of pages."""
        writer = PdfWriter()
        for i in range(pages):
            # Create a simple page (in a real scenario, you'd use a proper PDF library to create content)
            # For now, we'll create empty pages
            writer.add_blank_page(width=612, height=792)  # Standard letter size

        # Write to a temporary file
        temp_file = tempfile.NamedTemporaryFile(suffix='.pdf', delete=False)
        with open(temp_file.name, 'wb') as f:
            writer.write(f)
        return temp_file.name

    def test_rotate_pdf_pages_single_page(self):
        """Test rotating a single page in a PDF."""
        # Create a test PDF with 3 pages
        input_pdf = self.create_test_pdf(pages=3)

        try:
            # Create output file
            output_pdf = tempfile.NamedTemporaryFile(suffix='.pdf', delete=False).name

            # Rotate page 2 by 90 degrees
            rotate_pdf_pages(input_pdf, output_pdf, [2], 90)

            # Verify the output file exists
            assert os.path.exists(output_pdf)

            # Verify the PDF has the correct number of pages
            with open(output_pdf, 'rb') as f:
                reader = PdfReader(f)
                assert len(reader.pages) == 3

            # Clean up
            os.unlink(output_pdf)

        finally:
            os.unlink(input_pdf)

    def test_rotate_pdf_pages_multiple_pages(self):
        """Test rotating multiple pages in a PDF."""
        # Create a test PDF with 4 pages
        input_pdf = self.create_test_pdf(pages=4)

        try:
            # Create output file
            output_pdf = tempfile.NamedTemporaryFile(suffix='.pdf', delete=False).name

            # Rotate pages 1 and 3 by 180 degrees
            rotate_pdf_pages(input_pdf, output_pdf, [1, 3], 180)

            # Verify the output file exists
            assert os.path.exists(output_pdf)

            # Verify the PDF has the correct number of pages
            with open(output_pdf, 'rb') as f:
                reader = PdfReader(f)
                assert len(reader.pages) == 4

            # Clean up
            os.unlink(output_pdf)

        finally:
            os.unlink(input_pdf)

    def test_rotate_pdf_pages_invalid_page(self):
        """Test rotating a page number that doesn't exist."""
        # Create a test PDF with 2 pages
        input_pdf = self.create_test_pdf(pages=2)

        try:
            # Create output file
            output_pdf = tempfile.NamedTemporaryFile(suffix='.pdf', delete=False).name

            # Try to rotate page 5 (which doesn't exist)
            rotate_pdf_pages(input_pdf, output_pdf, [5], 90)

            # Verify the output file exists (should still be created)
            assert os.path.exists(output_pdf)

            # Verify the PDF has the correct number of pages
            with open(output_pdf, 'rb') as f:
                reader = PdfReader(f)
                assert len(reader.pages) == 2  # Should have same number of pages

            # Clean up
            os.unlink(output_pdf)

        finally:
            os.unlink(input_pdf)

    def test_merge_pdfs(self):
        """Test merging two PDF files."""
        # Create two test PDFs
        pdf1 = self.create_test_pdf(pages=2)
        pdf2 = self.create_test_pdf(pages=3)

        try:
            # Create output file
            output_pdf = tempfile.NamedTemporaryFile(suffix='.pdf', delete=False).name

            # Merge the PDFs
            merge_pdfs(pdf1, pdf2, output_pdf)

            # Verify the output file exists
            assert os.path.exists(output_pdf)

            # Verify the merged PDF has the correct total number of pages
            with open(output_pdf, 'rb') as f:
                reader = PdfReader(f)
                assert len(reader.pages) == 5  # 2 + 3 pages

            # Clean up
            os.unlink(output_pdf)

        finally:
            os.unlink(pdf1)
            os.unlink(pdf2)

    def test_merge_pdfs_empty_files(self):
        """Test merging PDFs with different page counts."""
        # Create two PDFs with different page counts
        pdf1 = self.create_test_pdf(pages=1)
        pdf2 = self.create_test_pdf(pages=4)

        try:
            # Create output file
            output_pdf = tempfile.NamedTemporaryFile(suffix='.pdf', delete=False).name

            # Merge the PDFs
            merge_pdfs(pdf1, pdf2, output_pdf)

            # Verify the output file exists
            assert os.path.exists(output_pdf)

            # Verify the merged PDF has the correct total number of pages
            with open(output_pdf, 'rb') as f:
                reader = PdfReader(f)
                assert len(reader.pages) == 5  # 1 + 4 pages

            # Clean up
            os.unlink(output_pdf)

        finally:
            os.unlink(pdf1)
            os.unlink(pdf2)

    def test_rotate_different_angles(self):
        """Test rotating pages with different angles."""
        # Create a test PDF
        input_pdf = self.create_test_pdf(pages=2)

        try:
            # Test different angles
            for angle in [90, 180, 270]:
                output_pdf = tempfile.NamedTemporaryFile(suffix='.pdf', delete=False).name

                rotate_pdf_pages(input_pdf, output_pdf, [1], angle)

                # Verify the output file exists and has correct page count
                assert os.path.exists(output_pdf)
                with open(output_pdf, 'rb') as f:
                    reader = PdfReader(f)
                    assert len(reader.pages) == 2

                os.unlink(output_pdf)

        finally:
            os.unlink(input_pdf)


    def test_convert_pdf_to_docx(self):
        """Test converting a PDF to DOCX format."""
        # Create a test PDF
        input_pdf = self.create_test_pdf(pages=2)

        try:
            # Create output file
            output_docx = tempfile.NamedTemporaryFile(suffix='.docx', delete=False).name

            # Convert PDF to DOCX
            convert_pdf_to_docx(input_pdf, output_docx)

            # Verify the output file exists
            assert os.path.exists(output_docx)

            # Verify it's a DOCX file (check file size > 0)
            assert os.path.getsize(output_docx) > 0

            # Clean up
            os.unlink(output_docx)

        finally:
            os.unlink(input_pdf)


    def test_order_pdf_pages(self):
        """Test reordering pages in a PDF."""
        # Create a test PDF with 6 pages
        input_pdf = self.create_test_pdf(pages=6)

        try:
            # Create output file
            output_pdf = tempfile.NamedTemporaryFile(suffix='.pdf', delete=False).name

            # Reorder first 6 pages as 1,3,2,5,4,6
            order_pdf_pages(input_pdf, output_pdf, 6, '1,3,2,5,4,6')

            # Verify the output file exists
            assert os.path.exists(output_pdf)

            # Verify the PDF has the correct number of pages
            with open(output_pdf, 'rb') as f:
                reader = PdfReader(f)
                assert len(reader.pages) == 6

            # Clean up
            os.unlink(output_pdf)

        finally:
            os.unlink(input_pdf)

    def test_order_pdf_pages_partial(self):
        """Test reordering only some pages in a PDF."""
        input_pdf = self.create_test_pdf(pages=8)
        try:
            output_pdf = tempfile.NamedTemporaryFile(suffix='.pdf', delete=False).name
            order_pdf_pages(input_pdf, output_pdf, 4, '4,1,3,2')
            assert os.path.exists(output_pdf)
            with open(output_pdf, 'rb') as f:
                reader = PdfReader(f)
                assert len(reader.pages) == 8
            os.unlink(output_pdf)
        finally:
            os.unlink(input_pdf)

    def test_rotate_verifies_rotation_metadata(self):
        """Verify rotated pages have the correct rotation value in PDF metadata."""
        input_pdf = self.create_test_pdf(pages=3)
        output_pdf = tempfile.NamedTemporaryFile(suffix='.pdf', delete=False).name
        try:
            rotate_pdf_pages(input_pdf, output_pdf, [1, 3], 90)
            with open(output_pdf, 'rb') as f:
                reader = PdfReader(f)
                assert reader.pages[0].rotation == 90
                assert reader.pages[1].rotation == 0
                assert reader.pages[2].rotation == 90
            os.unlink(output_pdf)
        finally:
            os.unlink(input_pdf)

    def test_rotate_270_degrees_metadata(self):
        """Verify 270-degree rotation is stored correctly in PDF metadata."""
        input_pdf = self.create_test_pdf(pages=2)
        output_pdf = tempfile.NamedTemporaryFile(suffix='.pdf', delete=False).name
        try:
            rotate_pdf_pages(input_pdf, output_pdf, [2], 270)
            with open(output_pdf, 'rb') as f:
                reader = PdfReader(f)
                assert reader.pages[0].rotation == 0
                assert reader.pages[1].rotation == 270
            os.unlink(output_pdf)
        finally:
            os.unlink(input_pdf)

    def test_order_raises_when_num_pages_exceeds_total(self):
        """order_pdf_pages raises ValueError when num_pages > total pages."""
        input_pdf = self.create_test_pdf(pages=3)
        output_pdf = tempfile.NamedTemporaryFile(suffix='.pdf', delete=False).name
        try:
            with pytest.raises(ValueError, match="only has"):
                order_pdf_pages(input_pdf, output_pdf, 5, '1,2,3,4,5')
        finally:
            os.unlink(input_pdf)
            if os.path.exists(output_pdf):
                os.unlink(output_pdf)

    def test_order_raises_when_order_has_wrong_count(self):
        """order_pdf_pages raises ValueError when new_order length != num_pages."""
        input_pdf = self.create_test_pdf(pages=4)
        output_pdf = tempfile.NamedTemporaryFile(suffix='.pdf', delete=False).name
        try:
            with pytest.raises(ValueError, match="exactly"):
                order_pdf_pages(input_pdf, output_pdf, 4, '1,2,3')
        finally:
            os.unlink(input_pdf)
            if os.path.exists(output_pdf):
                os.unlink(output_pdf)

    def test_order_raises_when_order_has_duplicates(self):
        """order_pdf_pages raises ValueError when new_order contains duplicate pages."""
        input_pdf = self.create_test_pdf(pages=4)
        output_pdf = tempfile.NamedTemporaryFile(suffix='.pdf', delete=False).name
        try:
            with pytest.raises(ValueError, match="exactly once"):
                order_pdf_pages(input_pdf, output_pdf, 4, '1,1,3,4')
        finally:
            os.unlink(input_pdf)
            if os.path.exists(output_pdf):
                os.unlink(output_pdf)

    def test_order_raises_when_order_not_integers(self):
        """order_pdf_pages raises ValueError when new_order contains non-integers."""
        input_pdf = self.create_test_pdf(pages=3)
        output_pdf = tempfile.NamedTemporaryFile(suffix='.pdf', delete=False).name
        try:
            with pytest.raises(ValueError, match="Invalid new_order"):
                order_pdf_pages(input_pdf, output_pdf, 3, 'a,b,c')
        finally:
            os.unlink(input_pdf)
            if os.path.exists(output_pdf):
                os.unlink(output_pdf)

    def test_convert_output_is_valid_docx(self):
        """Verify convert_pdf_to_docx produces a valid, openable DOCX file."""
        from docx import Document
        input_pdf = self.create_test_pdf(pages=1)
        output_docx = tempfile.NamedTemporaryFile(suffix='.docx', delete=False).name
        try:
            convert_pdf_to_docx(input_pdf, output_docx)
            doc = Document(output_docx)
            assert doc is not None
            os.unlink(output_docx)
        finally:
            os.unlink(input_pdf)

    def test_merge_preserves_all_pages_accessible(self):
        """Verify every page of a merged PDF is accessible."""
        pdf1 = self.create_test_pdf(pages=3)
        pdf2 = self.create_test_pdf(pages=2)
        output_pdf = tempfile.NamedTemporaryFile(suffix='.pdf', delete=False).name
        try:
            merge_pdfs(pdf1, pdf2, output_pdf)
            with open(output_pdf, 'rb') as f:
                reader = PdfReader(f)
                assert len(reader.pages) == 5
                for page in reader.pages:
                    assert page is not None
            os.unlink(output_pdf)
        finally:
            os.unlink(pdf1)
            os.unlink(pdf2)


class TestConvertPageFailures:
    """pdf2docx page failures must never cost content silently."""

    def make_pdf(self, path, text='Some text'):
        doc = fitz.open()
        page = doc.new_page()
        page.insert_text((72, 72), text, fontname='helv', fontsize=14)
        doc.save(path)
        doc.close()

    def fake_converter(self, monkeypatch, outcomes):
        """
        Replaces pdf2docx's Converter with a stub driving `outcomes`.

        Each outcome is either an exception to raise or None to succeed by
        writing a minimal DOCX. Records the kwargs of every convert() call.
        """
        calls = []

        class FakeConverter:
            def __init__(self, input_path):
                self.input_path = input_path

            def convert(self, output_path, **kwargs):
                calls.append(kwargs)
                outcome = outcomes[len(calls) - 1]
                if outcome is not None:
                    raise outcome
                Document().save(output_path)

            def close(self):
                pass

        monkeypatch.setattr(pdf_operations, 'Converter', FakeConverter)
        return calls

    def test_page_errors_are_not_ignored(self, monkeypatch, tmp_path):
        """
        pdf2docx's default drops a failed page and still reports success.

        Passing ignore_page_error=False is what makes the loss visible, so the
        retry and fallback below can act on it.
        """
        calls = self.fake_converter(monkeypatch, [None])
        input_pdf = tmp_path / 'in.pdf'
        self.make_pdf(str(input_pdf))

        convert_pdf_to_docx(str(input_pdf), str(tmp_path / 'out.docx'))

        assert calls[0]['ignore_page_error'] is False

    def test_retries_with_relaxed_layout_settings(self, monkeypatch, tmp_path):
        """A page that trips the default layout analysis gets a second chance."""
        calls = self.fake_converter(
            monkeypatch, [RuntimeError('Failed to merge docx_cell'), None]
        )
        input_pdf = tmp_path / 'in.pdf'
        output_docx = tmp_path / 'out.docx'
        self.make_pdf(str(input_pdf))

        convert_pdf_to_docx(str(input_pdf), str(output_docx))

        assert len(calls) == 2
        assert 'min_border_clearance' not in calls[0]
        assert calls[1]['min_border_clearance'] == 5.0
        assert output_docx.exists()

    def test_falls_back_to_text_when_every_attempt_fails(self, monkeypatch, tmp_path):
        """
        Layout is expendable; text is not.

        When no pdf2docx attempt produces a document, the fallback extraction
        has to run so the output still carries the PDF's text.
        """
        self.fake_converter(
            monkeypatch,
            [RuntimeError('first'), RuntimeError('second')],
        )
        input_pdf = tmp_path / 'in.pdf'
        output_docx = tmp_path / 'out.docx'
        self.make_pdf(str(input_pdf), text='Marker text to preserve')

        convert_pdf_to_docx(str(input_pdf), str(output_docx))

        text = '\n'.join(p.text for p in Document(str(output_docx)).paragraphs)
        assert 'Marker text to preserve' in text

    def test_a_clean_conversion_is_not_retried(self, monkeypatch, tmp_path):
        """The happy path must stay a single pass over the document."""
        calls = self.fake_converter(monkeypatch, [None])
        input_pdf = tmp_path / 'in.pdf'
        self.make_pdf(str(input_pdf))

        convert_pdf_to_docx(str(input_pdf), str(tmp_path / 'out.docx'))

        assert len(calls) == 1

    def test_coverage_is_total_when_the_text_all_arrives(self, tmp_path):
        input_pdf = tmp_path / 'in.pdf'
        output_docx = tmp_path / 'out.docx'
        self.make_pdf(str(input_pdf), text='Alpha beta gamma delta')

        document = Document()
        document.add_paragraph('Alpha beta gamma delta')
        document.save(str(output_docx))

        assert pdf_operations._text_coverage(str(input_pdf), str(output_docx)) == 1.0

    def test_coverage_notices_missing_text(self, tmp_path):
        input_pdf = tmp_path / 'in.pdf'
        output_docx = tmp_path / 'out.docx'
        self.make_pdf(str(input_pdf), text='Alpha beta gamma delta')

        Document().save(str(output_docx))

        assert pdf_operations._text_coverage(str(input_pdf), str(output_docx)) == 0.0

    def test_incomplete_output_is_reported(self, monkeypatch, tmp_path, capsys):
        """Dropped content must be announced, not left for the user to find."""
        self.fake_converter(monkeypatch, [None])  # writes an empty DOCX
        input_pdf = tmp_path / 'in.pdf'
        self.make_pdf(str(input_pdf), text='Content that will go missing')

        convert_pdf_to_docx(str(input_pdf), str(tmp_path / 'out.docx'))

        assert 'missing from the converted document' in capsys.readouterr().out

    def test_a_complete_conversion_is_quiet(self, monkeypatch, tmp_path, capsys):
        """The warning must not cry wolf on a conversion that lost nothing."""
        input_pdf = tmp_path / 'in.pdf'
        self.make_pdf(str(input_pdf), text='Everything here survives')

        class FullConverter:
            def __init__(self, input_path):
                pass

            def convert(self, output_path, **kwargs):
                document = Document()
                document.add_paragraph('Everything here survives')
                document.save(output_path)

            def close(self):
                pass

        monkeypatch.setattr(pdf_operations, 'Converter', FullConverter)
        convert_pdf_to_docx(str(input_pdf), str(tmp_path / 'out.docx'))

        assert 'missing from the converted document' not in capsys.readouterr().out

    def test_a_broken_coverage_check_never_costs_the_conversion(
        self, monkeypatch, tmp_path
    ):
        """The estimate is a diagnostic; it must not be able to fail a convert."""
        self.fake_converter(monkeypatch, [None])
        monkeypatch.setattr(
            pdf_operations,
            '_text_coverage',
            lambda *_: (_ for _ in ()).throw(RuntimeError('boom')),
        )
        input_pdf = tmp_path / 'in.pdf'
        output_docx = tmp_path / 'out.docx'
        self.make_pdf(str(input_pdf))

        convert_pdf_to_docx(str(input_pdf), str(output_docx))

        assert output_docx.exists()


class TestConvertRTL:
    """Right-to-left handling in convert_pdf_to_docx."""

    def make_pdf(self, text, path):
        """Write a one-page PDF containing text, using a Hebrew-capable font."""
        doc = fitz.open()
        page = doc.new_page()
        if HEBREW_FONT:
            page.insert_font(fontname='rtlfont', fontfile=HEBREW_FONT)
            page.insert_text((72, 72), text, fontname='rtlfont', fontsize=14)
        else:
            page.insert_text((72, 72), text, fontname='helv', fontsize=14)
        doc.save(path)
        doc.close()

    def body_xml(self, docx_path):
        return Document(docx_path).element.body.xml

    def test_rejects_an_unknown_rtl_mode(self):
        input_pdf = tempfile.NamedTemporaryFile(suffix='.pdf', delete=False).name
        try:
            self.make_pdf('Hello', input_pdf)
            with pytest.raises(ValueError, match='rtl must be one of'):
                convert_pdf_to_docx(input_pdf, '/tmp/unused.docx', rtl='sideways')
        finally:
            os.unlink(input_pdf)

    def test_ltr_pdf_gets_no_rtl_markup(self):
        """Regression guard: left-to-right documents must be untouched."""
        input_pdf = tempfile.NamedTemporaryFile(suffix='.pdf', delete=False).name
        output_docx = tempfile.NamedTemporaryFile(suffix='.docx', delete=False).name
        try:
            self.make_pdf('Plain English text, nothing to repair.', input_pdf)
            convert_pdf_to_docx(input_pdf, output_docx)

            xml = self.body_xml(output_docx)
            assert '<w:rtl' not in xml
            assert '<w:bidi' not in xml
        finally:
            os.unlink(input_pdf)
            if os.path.exists(output_docx):
                os.unlink(output_docx)

    def test_ltr_output_is_identical_with_rtl_off(self, tmp_path):
        """
        Golden guard for the RTL gating: on a left-to-right document, 'auto'
        must produce exactly what 'off' does, part for part.

        Compares the archive members rather than the .docx bytes, since the
        zip container stamps each member with the time it was written.
        """
        auto_docx = tmp_path / 'auto.docx'
        off_docx = tmp_path / 'off.docx'
        convert_pdf_to_docx(str(FIXTURES / 'ltr.pdf'), str(auto_docx))
        convert_pdf_to_docx(str(FIXTURES / 'ltr.pdf'), str(off_docx), rtl='off')

        with zipfile.ZipFile(auto_docx) as auto, zipfile.ZipFile(off_docx) as off:
            assert auto.namelist() == off.namelist()
            for name in auto.namelist():
                assert auto.read(name) == off.read(name), name

    @needs_hebrew_font
    def test_hebrew_pdf_gets_rtl_markup(self):
        input_pdf = tempfile.NamedTemporaryFile(suffix='.pdf', delete=False).name
        output_docx = tempfile.NamedTemporaryFile(suffix='.docx', delete=False).name
        try:
            self.make_pdf('שלום עולם', input_pdf)
            convert_pdf_to_docx(input_pdf, output_docx)

            xml = self.body_xml(output_docx)
            assert '<w:rtl' in xml
            assert '<w:bidi' in xml
        finally:
            os.unlink(input_pdf)
            if os.path.exists(output_docx):
                os.unlink(output_docx)

    @needs_hebrew_font
    def test_rtl_off_skips_the_repair_pass(self):
        input_pdf = tempfile.NamedTemporaryFile(suffix='.pdf', delete=False).name
        output_docx = tempfile.NamedTemporaryFile(suffix='.docx', delete=False).name
        try:
            self.make_pdf('שלום עולם', input_pdf)
            convert_pdf_to_docx(input_pdf, output_docx, rtl='off')

            assert '<w:rtl' not in self.body_xml(output_docx)
        finally:
            os.unlink(input_pdf)
            if os.path.exists(output_docx):
                os.unlink(output_docx)

    def test_rtl_on_forces_the_repair_pass(self):
        """'on' applies direction markup even when detection would say no."""
        input_pdf = tempfile.NamedTemporaryFile(suffix='.pdf', delete=False).name
        output_docx = tempfile.NamedTemporaryFile(suffix='.docx', delete=False).name
        try:
            self.make_pdf('Latin only', input_pdf)
            convert_pdf_to_docx(input_pdf, output_docx, rtl='on')

            # Nothing to mark, but the pass must run without damaging the file.
            assert Document(output_docx) is not None
        finally:
            os.unlink(input_pdf)
            if os.path.exists(output_docx):
                os.unlink(output_docx)


class TestConvertDocxToPdf:
    """DOCX to PDF conversion through LibreOffice."""

    def make_docx(self, path, text='Hello from a Word document'):
        document = Document()
        document.add_heading('Test Document', 0)
        document.add_paragraph(text)
        document.save(path)

    def fake_soffice(self, tmp_path, body):
        """
        Install a stand-in for the LibreOffice binary.

        Lets the surrounding plumbing — argument construction, output
        relocation, failure handling — be tested on machines without
        LibreOffice installed.
        """
        script = tmp_path / 'soffice'
        script.write_text('#!/usr/bin/env python3\n' + body)
        script.chmod(0o755)
        return str(script)

    def test_relocates_output_from_libreoffice_s_naming(self, tmp_path, monkeypatch):
        """LibreOffice writes <stem>.pdf into --outdir; we must move it."""
        stub = self.fake_soffice(tmp_path, '''
import sys, pathlib
outdir = pathlib.Path(sys.argv[sys.argv.index("--outdir") + 1])
stem = pathlib.Path(sys.argv[-1]).stem
(outdir / (stem + ".pdf")).write_bytes(b"%PDF-1.4 stub")
''')
        monkeypatch.setenv('SOFFICE_BIN', stub)

        input_docx = tmp_path / 'source.docx'
        self.make_docx(str(input_docx))
        output_pdf = tmp_path / 'a-different-name.pdf'

        convert_docx_to_pdf(str(input_docx), str(output_pdf))

        assert output_pdf.read_bytes() == b'%PDF-1.4 stub'

    def test_raises_when_libreoffice_exits_zero_without_output(self, tmp_path, monkeypatch):
        """LibreOffice can report success having produced nothing at all."""
        stub = self.fake_soffice(tmp_path, 'import sys; sys.exit(0)\n')
        monkeypatch.setenv('SOFFICE_BIN', stub)

        input_docx = tmp_path / 'source.docx'
        self.make_docx(str(input_docx))

        with pytest.raises(RuntimeError, match='produced no PDF'):
            convert_docx_to_pdf(str(input_docx), str(tmp_path / 'out.pdf'))

    def test_raises_with_stderr_on_failure(self, tmp_path, monkeypatch):
        stub = self.fake_soffice(
            tmp_path,
            'import sys; print("conversion blew up", file=sys.stderr); sys.exit(3)\n',
        )
        monkeypatch.setenv('SOFFICE_BIN', stub)

        input_docx = tmp_path / 'source.docx'
        self.make_docx(str(input_docx))

        with pytest.raises(RuntimeError, match='conversion blew up'):
            convert_docx_to_pdf(str(input_docx), str(tmp_path / 'out.pdf'))

    def test_raises_on_timeout(self, tmp_path, monkeypatch):
        stub = self.fake_soffice(tmp_path, 'import time; time.sleep(30)\n')
        monkeypatch.setenv('SOFFICE_BIN', stub)

        input_docx = tmp_path / 'source.docx'
        self.make_docx(str(input_docx))

        with pytest.raises(RuntimeError, match='timed out'):
            convert_docx_to_pdf(str(input_docx), str(tmp_path / 'out.pdf'), timeout=1)

    def test_uses_a_private_profile_directory(self, tmp_path, monkeypatch):
        """Concurrent conversions must not share a LibreOffice user profile."""
        stub = self.fake_soffice(tmp_path, '''
import sys, pathlib
outdir = pathlib.Path(sys.argv[sys.argv.index("--outdir") + 1])
(outdir / "argv.txt").write_text("\\n".join(sys.argv))
stem = pathlib.Path(sys.argv[-1]).stem
(outdir / (stem + ".pdf")).write_bytes(b"%PDF-1.4 stub")
''')
        monkeypatch.setenv('SOFFICE_BIN', stub)

        recorded = {}
        real_move = pdf_operations.shutil.move

        def capture(src, dst):
            recorded['argv'] = (Path(src).parent / 'argv.txt').read_text()
            return real_move(src, dst)

        monkeypatch.setattr(pdf_operations.shutil, 'move', capture)

        input_docx = tmp_path / 'source.docx'
        self.make_docx(str(input_docx))
        convert_docx_to_pdf(str(input_docx), str(tmp_path / 'out.pdf'))

        assert '-env:UserInstallation=file://' in recorded['argv']
        assert '--headless' in recorded['argv']

    def test_missing_libreoffice_explains_how_to_install_it(self, tmp_path, monkeypatch):
        monkeypatch.delenv('SOFFICE_BIN', raising=False)
        monkeypatch.setattr(pdf_operations.shutil, 'which', lambda _: None)

        input_docx = tmp_path / 'source.docx'
        self.make_docx(str(input_docx))

        with pytest.raises(RuntimeError, match='libreoffice-writer'):
            convert_docx_to_pdf(str(input_docx), str(tmp_path / 'out.pdf'))

    @needs_soffice
    def test_converts_a_docx_to_a_readable_pdf(self):
        input_docx = tempfile.NamedTemporaryFile(suffix='.docx', delete=False).name
        output_pdf = tempfile.NamedTemporaryFile(suffix='.pdf', delete=False).name
        try:
            self.make_docx(input_docx)
            convert_docx_to_pdf(input_docx, output_pdf)

            assert os.path.exists(output_pdf)
            reader = PdfReader(output_pdf)
            assert len(reader.pages) >= 1
            assert 'Hello from a Word document' in reader.pages[0].extract_text()
        finally:
            os.unlink(input_docx)
            if os.path.exists(output_pdf):
                os.unlink(output_pdf)

    @needs_soffice
    def test_writes_to_the_requested_path_not_libreoffice_s_choice(self):
        """LibreOffice names output after the input stem; we must override it."""
        input_docx = tempfile.NamedTemporaryFile(suffix='.docx', delete=False).name
        output_pdf = os.path.join(tempfile.mkdtemp(), 'a-different-name.pdf')
        try:
            self.make_docx(input_docx)
            convert_docx_to_pdf(input_docx, output_pdf)

            assert os.path.exists(output_pdf)
        finally:
            os.unlink(input_docx)
            if os.path.exists(output_pdf):
                os.unlink(output_pdf)

    def test_rejects_a_file_that_is_not_a_document(self, tmp_path):
        """
        LibreOffice would happily render these bytes as plain text and report
        success, so the check has to happen before it is invoked.
        """
        bogus = tmp_path / 'bogus.docx'
        bogus.write_bytes(b'this is not a docx at all')
        output = tmp_path / 'out.pdf'

        with pytest.raises(ValueError, match='not a Word document'):
            convert_docx_to_pdf(str(bogus), str(output))
        assert not output.exists()

    def test_rejects_a_zip_that_is_not_a_document(self, tmp_path):
        archive = tmp_path / 'archive.docx'
        with zipfile.ZipFile(archive, 'w') as zf:
            zf.writestr('hello.txt', 'not a document')
        output = tmp_path / 'out.pdf'

        with pytest.raises(ValueError, match='no word/document.xml'):
            convert_docx_to_pdf(str(archive), str(output))
        assert not output.exists()


if __name__ == "__main__":
    pytest.main([__file__])

def make_text_pdf(path: Path, lines: list[str]) -> None:
    """Writes a one-page PDF with large Helvetica text and identifying metadata."""
    document = fitz.open()
    page = document.new_page()
    for index, line in enumerate(lines):
        page.insert_text((72, 120 + index * 60), line, fontsize=28, fontname='helv')
    document.set_metadata({
        'author': 'Secret Author',
        'creator': 'Tracking Creator',
        'producer': 'Tracking Producer',
        'title': 'Confidential Title',
    })
    document.save(path)
    document.close()


class TestSanitizePdf:
    @needs_tesseract
    def test_sanitize_rebuilds_searchable_pdf_without_metadata(self, tmp_path):
        source = tmp_path / 'source.pdf'
        output = tmp_path / 'sanitized.pdf'
        make_text_pdf(source, ['Quarterly Report', 'Sanitize Me Please'])

        sanitize_pdf(str(source), str(output), lang='eng')

        reader = PdfReader(str(output))
        assert len(reader.pages) == 1
        assert all(value == '' for value in reader.metadata.values())

        text = reader.pages[0].extract_text()
        assert 'Quarterly' in text
        assert 'Sanitize' in text

        with fitz.open(output) as rebuilt:
            fonts = {font[3] for font in rebuilt[0].get_fonts()}
        assert not any('Helvetica' in name for name in fonts)

    @needs_tesseract
    def test_sanitize_keeps_page_count(self, tmp_path):
        source = tmp_path / 'source.pdf'
        output = tmp_path / 'sanitized.pdf'
        document = fitz.open()
        for number in range(3):
            document.new_page().insert_text((72, 120), f'Page {number + 1}', fontsize=28)
        document.save(source)
        document.close()

        sanitize_pdf(str(source), str(output), lang='eng', dpi=150)

        assert len(PdfReader(str(output)).pages) == 3

    def test_sanitize_rejects_malformed_lang(self, tmp_path):
        with pytest.raises(ValueError, match='Invalid OCR language'):
            sanitize_pdf(str(FIXTURES / 'ltr.pdf'), str(tmp_path / 'out.pdf'), lang='eng; rm -rf')

    def test_missing_language_pack_is_reported(self):
        with pytest.raises(ValueError, match='heb'):
            pdf_operations._check_ocr_langs('eng+heb', ['eng', 'osd'])

    def test_installed_languages_pass(self):
        pdf_operations._check_ocr_langs('eng+heb', ['eng', 'heb'])

    def test_missing_tesseract_raises_runtime_error(self, tmp_path, monkeypatch):
        monkeypatch.delenv('TESSERACT_CMD', raising=False)
        monkeypatch.setattr(pdf_operations.shutil, 'which', lambda name: None)
        with pytest.raises(RuntimeError, match='Tesseract is required'):
            sanitize_pdf(str(FIXTURES / 'ltr.pdf'), str(tmp_path / 'out.pdf'))


def make_dirty_pdf(path: Path) -> None:
    """Writes a PDF carrying each kind of hidden data scrub_pdf should strip."""
    document = fitz.open()
    page = document.new_page()
    page.insert_text((72, 100), 'Visible text', fontsize=20)
    page.insert_text((72, 200), 'Invisible secret', fontsize=20, render_mode=3)
    page.add_text_annot((300, 300), 'Reviewer comment')
    page.insert_link({'kind': fitz.LINK_URI, 'from': fitz.Rect(72, 80, 200, 110), 'uri': 'https://tracker.example'})
    document.embfile_add('attachment.txt', b'attached secret')
    document.set_metadata({'author': 'Secret Author', 'creator': 'Tracking Creator'})
    document.set_xml_metadata('<x:xmpmeta xmlns:x="adobe:ns:meta/">Secret XMP</x:xmpmeta>')
    document.save(path)
    document.close()


def make_clipped_edit_pdf(tmp_path: Path) -> Path:
    """
    Mimics an editor deleting a line: the page is drawn twice through clipped
    form XObjects that skip the line and close the gap, but the line's text
    is still in the file.
    """
    original = tmp_path / 'original.pdf'
    source = fitz.open()
    page = source.new_page()
    page.insert_text((72, 100), 'Kept first line', fontsize=14)
    page.insert_text((72, 130), 'Deleted secret line', fontsize=14)
    page.insert_text((72, 160), 'Kept last line', fontsize=14)
    source.save(original)
    source.close()

    edited = tmp_path / 'edited.pdf'
    with fitz.open(original) as source:
        document = fitz.open()
        page = document.new_page()
        width = page.rect.width
        page.show_pdf_page(fitz.Rect(0, 0, width, 110), source, 0, clip=fitz.Rect(0, 0, width, 110))
        page.show_pdf_page(fitz.Rect(0, 110, width, 812), source, 0, clip=fitz.Rect(0, 140, width, 842))
        document.save(edited)
        document.close()
    return edited


def drawn_text(path) -> str:
    """Every character drawn on page 1, clipped or not, in drawing order."""
    with fitz.open(path) as document:
        return ''.join(chr(c[0]) for s in document[0].get_texttrace() for c in s['chars'])


class TestScrubPdf:
    def test_scrub_strips_hidden_data_and_keeps_visible_text(self, tmp_path):
        source = tmp_path / 'source.pdf'
        output = tmp_path / 'scrubbed.pdf'
        make_dirty_pdf(source)

        scrub_pdf(str(source), str(output))

        with fitz.open(output) as scrubbed:
            page = scrubbed[0]
            assert page.get_text().strip() == 'Visible text'
            assert list(page.annots()) == []
            assert page.get_links() == []
            assert scrubbed.embfile_count() == 0
            assert scrubbed.get_xml_metadata() == ''
            assert all(not value for key, value in scrubbed.metadata.items() if key != 'format')

        raw = output.read_bytes()
        for secret in (b'Invisible secret', b'Reviewer comment', b'tracker.example', b'Secret Author', b'Secret XMP'):
            assert secret not in raw

    def test_scrub_keeps_page_count_and_real_text(self, tmp_path):
        output = tmp_path / 'scrubbed.pdf'

        scrub_pdf(str(FIXTURES / 'mixed.pdf'), str(output))

        with fitz.open(FIXTURES / 'mixed.pdf') as original, fitz.open(output) as scrubbed:
            assert scrubbed.page_count == original.page_count
            assert scrubbed[0].get_text() == original[0].get_text()

    def test_scrub_removes_form_fields(self, tmp_path):
        source = tmp_path / 'form.pdf'
        output = tmp_path / 'scrubbed.pdf'
        document = fitz.open()
        page = document.new_page()
        widget = fitz.Widget()
        widget.field_type = fitz.PDF_WIDGET_TYPE_TEXT
        widget.field_name = 'ssn'
        widget.field_value = '123-45-6789'
        widget.rect = fitz.Rect(72, 72, 300, 100)
        page.add_widget(widget)
        document.save(source)
        document.close()

        scrub_pdf(str(source), str(output))

        with fitz.open(output) as scrubbed:
            assert list(scrubbed[0].widgets()) == []
        assert b'123-45-6789' not in output.read_bytes()

    def test_scrub_removes_text_deleted_by_clipping(self, tmp_path):
        """
        Editors often "delete" a region by redrawing the page twice through
        clipped form XObjects that skip it. The text stays in the file and
        copy-paste or an ATS parser still reads it.
        """
        edited = make_clipped_edit_pdf(tmp_path)
        assert 'Deleted secret line' in drawn_text(edited)

        output = tmp_path / 'scrubbed.pdf'
        scrub_pdf(str(edited), str(output))

        drawn = drawn_text(output)
        assert 'Deleted secret line' not in drawn
        assert drawn.count('Kept first line') == 1
        assert drawn.count('Kept last line') == 1
        assert 'Deleted' not in PdfReader(str(output)).pages[0].extract_text()

    def test_scrub_removes_transparent_text(self, tmp_path):
        source = tmp_path / 'source.pdf'
        output = tmp_path / 'scrubbed.pdf'
        document = fitz.open()
        page = document.new_page()
        page.insert_text((72, 100), 'Visible text', fontsize=20)
        page.insert_text((72, 200), 'Transparent keywords', fontsize=20, fill_opacity=0)
        document.save(source)
        document.close()

        scrub_pdf(str(source), str(output))

        with fitz.open(output) as scrubbed:
            assert scrubbed[0].get_text().strip() == 'Visible text'

    def test_scrub_rejects_non_pdf(self, tmp_path):
        bogus = tmp_path / 'bogus.pdf'
        bogus.write_bytes(b'not a pdf')
        with pytest.raises(ValueError):
            scrub_pdf(str(bogus), str(tmp_path / 'out.pdf'))


def render_difference(first, second) -> float:
    """Share of pixels on page 1 that differ clearly between two PDFs."""
    from PIL import Image, ImageChops
    images = []
    for path in (first, second):
        with fitz.open(path) as document:
            images.append(Image.open(io.BytesIO(document[0].get_pixmap(dpi=100).tobytes('png'))).convert('L'))
    histogram = ImageChops.difference(*images).histogram()
    return sum(histogram[65:]) / (images[0].width * images[0].height)


class TestAtsPdf:
    def test_text_layer_follows_reading_order(self, tmp_path):
        """The file draws the lower line first; extraction must still read top down."""
        source = tmp_path / 'source.pdf'
        output = tmp_path / 'ats.pdf'
        document = fitz.open()
        page = document.new_page()
        page.insert_text((72, 300), 'Second line drawn first', fontsize=14)
        page.insert_text((72, 100), 'First line drawn last', fontsize=14)
        document.save(source)
        document.close()

        # pypdf reads in drawing order, like copy-paste and many ATS parsers.
        assert PdfReader(str(source)).pages[0].extract_text().startswith('Second')

        ats_pdf(str(source), str(output))

        lines = PdfReader(str(output)).pages[0].extract_text().splitlines()
        assert lines == ['First line drawn last', 'Second line drawn first']

    def test_removes_hidden_text_and_the_original_fonts(self, tmp_path):
        source = tmp_path / 'source.pdf'
        output = tmp_path / 'ats.pdf'
        make_dirty_pdf(source)

        ats_pdf(str(source), str(output))

        assert PdfReader(str(output)).pages[0].extract_text().strip() == 'Visible text'
        with fitz.open(output) as rebuilt:
            fonts = rebuilt[0].get_fonts(full=True)
            assert [font[3] for font in fonts] == ['GlyphLessFont']
            # Embedded: PDF editors warn about a font they cannot find.
            assert rebuilt.extract_font(fonts[0][0])[3]
            assert all(not value for key, value in rebuilt.metadata.items() if key != 'format')

    def test_removes_text_deleted_by_clipping(self, tmp_path):
        edited = make_clipped_edit_pdf(tmp_path)
        output = tmp_path / 'ats.pdf'

        ats_pdf(str(edited), str(output))

        text = PdfReader(str(output)).pages[0].extract_text()
        assert 'Deleted' not in text
        assert text.splitlines() == ['Kept first line', 'Kept last line']

    def test_page_looks_the_same(self, tmp_path):
        output = tmp_path / 'ats.pdf'

        ats_pdf(str(FIXTURES / 'ltr.pdf'), str(output))

        # Outlines and font rendering anti-alias edges a little differently.
        assert render_difference(FIXTURES / 'ltr.pdf', output) < 0.01
        with fitz.open(FIXTURES / 'ltr.pdf') as original, fitz.open(output) as rebuilt:
            assert rebuilt[0].get_text() == original[0].get_text()

    def test_right_to_left_text_keeps_its_words(self, tmp_path):
        for name in ('hebrew.pdf', 'mixed.pdf', 'textbox.pdf'):
            output = tmp_path / name

            ats_pdf(str(FIXTURES / name), str(output))

            with fitz.open(FIXTURES / name) as original, fitz.open(output) as rebuilt:
                assert sorted(rebuilt[0].get_text().split()) == sorted(original[0].get_text().split()), name

    def test_glyphless_font_is_a_valid_truetype_font(self):
        from pdf_editor.page_text import _glyphless_truetype

        font = fitz.Font(fontbuffer=_glyphless_truetype())
        assert font.name.startswith('GlyphLessFont')
        assert font.glyph_count == 2

    def test_columns_are_read_one_after_the_other(self, tmp_path):
        source = tmp_path / 'columns.pdf'
        output = tmp_path / 'ats.pdf'
        filler = 'alpha beta gamma delta epsilon zeta eta theta iota kappa lambda mu nu xi omicron'
        document = fitz.open()
        page = document.new_page()
        page.insert_textbox(fitz.Rect(50, 100, 280, 200), 'LEFTONE ' + filler, fontsize=10)
        page.insert_textbox(fitz.Rect(320, 100, 550, 200), 'RIGHTONE ' + filler, fontsize=10)
        page.insert_textbox(fitz.Rect(50, 210, 280, 300), 'LEFTTWO ' + filler, fontsize=10)
        page.insert_textbox(fitz.Rect(320, 210, 550, 300), 'RIGHTTWO ' + filler, fontsize=10)
        document.save(source)
        document.close()

        ats_pdf(str(source), str(output))

        text = PdfReader(str(output)).pages[0].extract_text()
        markers = ['LEFTONE', 'LEFTTWO', 'RIGHTONE', 'RIGHTTWO']
        assert sorted(markers, key=text.index) == markers

    def test_dates_beside_titles_are_read_row_by_row(self, tmp_path):
        """A narrow column of dates beside job titles is a table, not two columns."""
        source = tmp_path / 'cv.pdf'
        output = tmp_path / 'ats.pdf'
        document = fitz.open()
        page = document.new_page()
        for top, (dates, title) in zip((100, 160), [('2020 - 2024', 'Senior Engineer'), ('2016 - 2020', 'Engineer')]):
            page.insert_text((50, top), dates, fontsize=9)
            page.insert_text((130, top), title, fontsize=11)
            page.insert_text((130, top + 15), f'Did things as {title}', fontsize=9)
        document.save(source)
        document.close()

        ats_pdf(str(source), str(output))

        text = PdfReader(str(output)).pages[0].extract_text()
        markers = ['2020 - 2024', 'Senior Engineer', '2016 - 2020', 'Did things as Engineer']
        assert sorted(markers, key=text.index) == markers

    @needs_hebrew_font
    def test_right_to_left_columns_are_read_right_first(self, tmp_path):
        source = tmp_path / 'columns.pdf'
        output = tmp_path / 'ats.pdf'
        filler = 'זהו טקסט ארוך בעברית שנועד למלא כמה שורות בתוך הטור כדי לבדוק את סדר הקריאה'
        document = fitz.open()
        page = document.new_page()
        page.insert_htmlbox(fitz.Rect(320, 100, 550, 260), f'<p dir="rtl">ימני {filler}</p>')
        page.insert_htmlbox(fitz.Rect(50, 100, 280, 260), f'<p dir="rtl">שמאלי {filler}</p>')
        document.save(source)
        document.close()

        ats_pdf(str(source), str(output))

        with fitz.open(output) as rebuilt:
            text = rebuilt[0].get_text()
        assert text.index('ימני') < text.index('שמאלי')

    @needs_hebrew_font
    def test_rotated_right_to_left_line_keeps_its_words(self, tmp_path):
        line = tmp_path / 'line.pdf'
        source = tmp_path / 'rotated.pdf'
        output = tmp_path / 'ats.pdf'
        document = fitz.open()
        page = document.new_page()
        page.insert_htmlbox(fitz.Rect(50, 50, 400, 100), '<p dir="rtl">שלום עולם טל: 03-9377218 סוף</p>')
        document.save(line)
        document.close()
        with fitz.open(line) as single:
            document = fitz.open()
            document.new_page().show_pdf_page(fitz.Rect(50, 50, 500, 700), single, 0, rotate=90)
            document.save(source)
            document.close()

        ats_pdf(str(source), str(output))

        with fitz.open(source) as original, fitz.open(output) as rebuilt:
            assert sorted(rebuilt[0].get_text().split()) == sorted(original[0].get_text().split())

    def test_lone_surrogate_from_a_broken_tounicode_map_does_not_crash(self, tmp_path):
        source = tmp_path / 'broken.pdf'
        output = tmp_path / 'ats.pdf'
        document = fitz.open()
        page = document.new_page()
        page.insert_text((72, 100), 'ABC Hello', fontsize=14)
        document.save(source)
        document.close()
        # Point the font's ToUnicode entry for 'A' at a lone high surrogate.
        with fitz.open(source) as document:
            font_xref = document[0].get_fonts()[0][0]
            cmap = document.get_new_xref()
            document.update_object(cmap, '<<>>')
            document.update_stream(cmap, (
                b'/CIDInit /ProcSet findresource begin 12 dict begin begincmap '
                b'/CMapName /Broken def 1 begincodespacerange <00> <FF> endcodespacerange '
                b'1 beginbfchar <41> <D800> endbfchar endcmap CMapName currentdict /CMap '
                b'defineresource pop end end'
            ))
            document.xref_set_key(font_xref, 'ToUnicode', f'{cmap} 0 R')
            document.save(tmp_path / 'broken2.pdf')

        ats_pdf(str(tmp_path / 'broken2.pdf'), str(output))

        with fitz.open(output) as rebuilt:
            text = rebuilt[0].get_text()
        assert 'Hello' in text and '�' in text
